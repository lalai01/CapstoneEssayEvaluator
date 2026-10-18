"""
Read a grading rubric from an uploaded PDF, Excel (.xlsx) or CSV file and
turn it into criteria for the rubric editor. Nothing is saved here: the user
reviews, edits and saves the result in the editor.

Rubrics come in many layouts: criteria as rows or as columns, numbered
levels (4, 3, 2, 1) or named ones (Excellent / Adequate / Needs Work),
bullet points in cells, group labels, comment rows. So the reading is done
by an AI, in order of preference:
  1. Gemini, which reads the PDF itself (layout and all) or the sheet's text;
  2. DeepSeek, from the extracted text;
  3. a built-in table reader for simple grid rubrics, when no AI is available.

Result: {"title", "criteria": [{"label", "max_score", "weight", "levels"}],
         "source", "warnings"} with levels {score: description}.
"""
import base64
import csv
import io
import json
import os
import re

import requests

MAX_LEVELS = 10            # levels kept per criterion
MIN_LEVELS = 2
MAX_POINTS = 100           # highest score a criterion can have (evaluator.MAX_SCORE_LIMIT)
MAX_DESCRIPTION = 600      # characters kept per level description

INSTRUCTIONS = """Extract the essay grading rubric from this document.

- Each criterion (trait, category, row or column that an essay is scored on)
  becomes one item in "criteria". Ignore title rows, group labels that only
  group other criteria, "comments"/"reviewer" rows, totals, and pages of
  standards that aren't part of the scoring grid.
- For each criterion list its performance levels with a numeric "score":
  use the rubric's own numbers when it has them (e.g. 4, 3, 2, 1). If the
  levels are names only (e.g. Excellent / Adequate / Needs Work), number
  them from 1 for the lowest upward.
- Keep the rubric's real points. A level with a points range (e.g.
  "Excellent (18-20)") is ONE level: "score" = the lowest points of the
  range (18), "max" = the highest (20), "label" = its name (Excellent).
  Never renumber point ranges to 1-4.
- "max_score": the criterion's maximum points (e.g. 20), if stated.
- Include every criterion row of the scoring grid, even when a sheet or page
  has many; don't merge or skip criteria.
- "description": the level's descriptor, copied faithfully; join bullet
  points with "; ". Leave it empty only if the cell is empty.
- "weight": the criterion's relative importance if the rubric states points,
  percentages or weights; otherwise 1.
- "title": the rubric's title, or a short descriptive one."""

SCHEMA = {
    "type": "OBJECT",
    "properties": {
        "title": {"type": "STRING"},
        "criteria": {
            "type": "ARRAY",
            "items": {
                "type": "OBJECT",
                "properties": {
                    "name": {"type": "STRING"},
                    "weight": {"type": "NUMBER"},
                    "max_score": {"type": "INTEGER"},
                    "levels": {
                        "type": "ARRAY",
                        "items": {
                            "type": "OBJECT",
                            "properties": {
                                "score": {"type": "INTEGER"},
                                "max": {"type": "INTEGER"},
                                "label": {"type": "STRING"},
                                "description": {"type": "STRING"},
                            },
                            "required": ["score", "description"],
                        },
                    },
                },
                "required": ["name", "levels"],
            },
        },
    },
    "required": ["title", "criteria"],
}


# ---------------------------------------------------------------- file → text / tables

def _clean(cell):
    return re.sub(r"\s+", " ", str(cell or "")).strip()


def pdf_tables(data):
    """[[row cells]] for the largest table on each page."""
    import fitz
    tables = []
    with fitz.open(stream=data, filetype="pdf") as doc:
        for page in doc:
            found = page.find_tables().tables
            if found:
                biggest = max(found, key=lambda t: t.row_count * t.col_count)
                tables.append([[_clean(c) for c in row] for row in biggest.extract()])
    return tables


def pdf_text(data):
    import fitz
    with fitz.open(stream=data, filetype="pdf") as doc:
        return "\n".join(page.get_text() for page in doc)


def sheet_tables(data, suffix):
    """[[row cells]] per sheet (.xlsx) or for the CSV."""
    if suffix == ".csv":
        text = data.decode("utf-8-sig", errors="replace")
        return [[[_clean(c) for c in row] for row in csv.reader(io.StringIO(text))]]
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    tables = []
    for ws in wb.worksheets:
        rows = [[_clean(c) for c in row] for row in ws.iter_rows(values_only=True)]
        rows = [r for r in rows if any(r)]
        if rows:
            tables.append(rows)
    return tables


def tables_as_text(tables):
    return "\n\n".join(
        "\n".join(" | ".join(row) for row in table) for table in tables
    )


# ---------------------------------------------------------------- AI readers

def _gemini(parts):
    import gemini_ocr
    if not gemini_ocr.enabled():
        raise RuntimeError("Gemini isn't configured")
    key = os.environ["GEMINI_API_KEY"].strip()
    body = {
        "contents": [{"parts": parts}],
        "generationConfig": {"temperature": 0, "responseMimeType": "application/json",
                             "responseSchema": SCHEMA},
    }
    errors = []
    for model in gemini_ocr._models():
        resp = requests.post(gemini_ocr.API_URL.format(model=model), json=body, timeout=120,
                             headers={"x-goog-api-key": key})
        if resp.status_code != 200:
            errors.append(f"{model}: HTTP {resp.status_code}")
            continue
        candidate = resp.json()["candidates"][0]
        return json.loads("".join(p.get("text", "") for p in candidate["content"]["parts"]))
    raise RuntimeError("; ".join(errors))


def _deepseek(text):
    from ai_models import call_deepseek
    if not os.environ.get("DEEPSEEK_API_KEY"):
        raise RuntimeError("DeepSeek isn't configured")
    shape = ('{"title": "...", "criteria": [{"name": "...", "weight": 1, "max_score": 20, '
             '"levels": [{"score": 18, "max": 20, "label": "Excellent", "description": "..."}]}]}')
    reply = call_deepseek(
        "You convert grading rubrics into JSON. Reply with a single JSON object only.",
        f"{INSTRUCTIONS}\n\nReply with JSON shaped like:\n{shape}\n\nRubric (table cells "
        f"separated by |):\n\n{text[:20000]}",
        json_mode=True, temperature=0.0,
    )["text"]
    return json.loads(reply)


# ---------------------------------------------------------------- built-in table reader

LEVEL_WORDS = re.compile(
    r"^(\d+|excellent|exemplary|advanced|exceeds|proficient|good|adequate|satisfactory|"
    r"developing|approaching|fair|basic|needs work|needs improvement|beginning|emerging|"
    r"poor|weak|unsatisfactory|incomplete|limited)\b", re.I)
SKIP_ROWS = re.compile(r"comment|reviewer|total|score", re.I)
# Header cells that name the column of criteria rather than a level.
HEADER_WORDS = re.compile(r"^(scores?|points?|levels?|criteri(a|on)|traits?|categor(y|ies))$", re.I)


def _is_level_label(cell):
    return bool(cell) and len(cell.split()) <= 4 and bool(LEVEL_WORDS.match(cell))


POINTS_RANGE = re.compile(r"(\d+)\s*(?:-|–|—|to)\s*(\d+)")
POINTS_TEXT = re.compile(r"[(\[]?\s*\d+\s*(?:(?:-|–|—|to)\s*\d+)?\s*(?:pts?\.?|points?)?\s*[)\]]?", re.I)


def _level_of(cell):
    """A header cell → {"score", "top", "label"}: "Excellent (18-20)" gives
    score 18, top 20; "4" gives 4, 4; "5 Mastery" gives 5, 5, "Mastery";
    a name alone gives no score (numbered later)."""
    label = _clean(POINTS_TEXT.sub(" ", cell))
    span = POINTS_RANGE.search(cell)
    if span:
        low, high = sorted((int(span.group(1)), int(span.group(2))))
        return {"score": low, "top": high, "label": label}
    number = re.search(r"\d+", cell)
    score = int(number.group()) if number else None
    return {"score": score, "top": score, "label": label}


def _header(row):
    """
    For a header row of level labels ("Score | 4 | 3 | 2 | 1 | NS" or
    "Criterion | Excellent (18-20) | Proficient (15-17) | ..."): its labelled
    columns as [(column, level or None)], None for columns that aren't levels
    (e.g. "NS", "Not Scored", "Maximum"). None if the row isn't a header.
    """
    cells = [(i, c) for i, c in enumerate(row) if i > 0 and c and not HEADER_WORDS.match(c)]
    levels = [(i, _level_of(c)) for i, c in cells if _is_level_label(c)]
    if len(levels) < MIN_LEVELS or len(cells) - len(levels) > 1:
        return None
    scores = [lv["score"] for _, lv in levels]
    if None in scores or len(set(scores)) != len(scores):
        # Named levels, best first: number them down to 1.
        for n, (_, lv) in zip(range(len(levels), 0, -1), levels):
            lv["score"] = lv["top"] = n
    level_at = dict(levels)
    return [(i, level_at.get(i)) for i, _ in cells]


def _row_levels(row, header):
    """[(level, description)] for a criterion row under the given header."""
    filled = [(i, c) for i, c in enumerate(row) if i > 0 and c]
    if header is None:
        # No header: the level columns count down from the left.
        n = len(row) - 1
        return [({"score": n - (i - 1), "top": n - (i - 1), "label": ""}, c) for i, c in filled]
    columns = dict(header)
    if all(i in columns for i, _ in filled):
        pairs = [(columns[i], c) for i, c in filled]       # cells line up with the header
    else:
        # Merged cells shift columns (common in PDFs): match by order instead.
        pairs = [(lv, c) for (_, lv), (_, c) in zip(header, filled)]
    return [(lv, c) for lv, c in pairs if lv is not None]


def table_rubric(table):
    """
    Criteria from a grid: one criterion per row, one level per column (left
    = best, as rubrics are usually laid out). A header row of level labels
    sets the scores, and may repeat further down with a different scale; a
    rotated criterion name that comes out as its own row after the
    descriptions is attached to them. Returns [] if nothing looks like a
    rubric row.
    """
    rows = [r for r in table if any(r)]
    if not rows:
        return []
    width = max(len(r) for r in rows)
    if width - 1 < MIN_LEVELS:
        return []
    rows = [r + [""] * (width - len(r)) for r in rows]
    criteria, header, unnamed, had_header = [], None, None, False
    for row in rows:
        found = _header(row)
        if found:
            header, unnamed, had_header = found, None, True
            continue
        if _is_level_label(row[0]):
            return []   # levels as rows: a sideways rubric this reader can't follow
        filled = [c for c in row if c]
        if len(filled) == 1 and unnamed and len(filled[0].split()) <= 6:
            unnamed["name"] = filled[0]
            criteria.append(unnamed)
            unnamed = None
            continue
        levels = _row_levels(row, header)
        if len(levels) < MIN_LEVELS:
            continue
        item = {"name": row[0], "weight": 1, "max_score": max(lv["top"] for lv, _ in levels), "levels": [
            {"score": lv["score"], "max": lv["top"], "label": lv["label"], "description": d} for lv, d in levels]}
        if not row[0]:
            unnamed = item
        elif not SKIP_ROWS.search(row[0]):
            criteria.append(item)
    # A rubric describes each level in sentences; a row of short labels (e.g.
    # grade-level headers) isn't a criterion. Without a score header, a
    # single row is too little evidence that the table is a rubric.
    def described(c):
        texts = [lv["description"] for lv in c["levels"] if lv["description"]]
        return texts and sum(map(len, texts)) / len(texts) >= 25
    criteria = [c for c in criteria if described(c)]
    return criteria if had_header or len(criteria) >= 2 else []


# ---------------------------------------------------------------- normalize

def _int(value, default):
    try:
        return int(value)
    except (TypeError, ValueError):
        return default


def _normalize(raw):
    """AI or table output → editor criteria; drops anything unusable."""
    criteria, warnings, seen = [], [], set()
    for item in raw:
        name = _clean(item.get("name")).rstrip(":")
        # Some PDFs store text twice: "Development: Analysis Development: Analysis".
        name = re.sub(r"^(.{6,}?)\s+\1\b.*$", r"\1", name)
        levels = [lv for lv in (item.get("levels") or []) if isinstance(lv, dict)]
        scored = []   # (lowest score, highest score, description)
        for lv in levels:
            try:
                low = int(lv.get("score"))
            except (TypeError, ValueError):
                continue
            high = max(low, _int(lv.get("max"), low))
            text, label = _clean(lv.get("description")), _clean(lv.get("label"))
            # Keep the level's name ("Excellent: ...") when it isn't just a number.
            if label and text and not label.isdigit() and not text.lower().startswith(label.lower()):
                text = f"{label}: {text}"
            scored.append((low, high, text[:MAX_DESCRIPTION]))
        if not name or len(scored) < MIN_LEVELS:
            continue
        distinct = sorted({s for s, _, _ in scored})
        top = max(h for _, h, _ in scored)
        if distinct[0] < 0 or top > MAX_POINTS or len(distinct) > MAX_LEVELS:
            # Out of range or one level per point: keep the order, renumber 1..n.
            if len(distinct) > MAX_LEVELS:
                warnings.append(f"“{name}” had more than {MAX_LEVELS} levels; only the top {MAX_LEVELS} were kept.")
                distinct = distinct[-MAX_LEVELS:]
            rank = {s: i + 1 for i, s in enumerate(distinct)}
            scored = [(rank[s], rank[s], d) for s, _, d in scored if s in rank]
        elif distinct[0] == 0:
            # Scores start at 0, but scoring starts at 1: a band "0-10" starts
            # at 1; plain 0/1/2 levels move up by one.
            if any(h > s for s, h, _ in scored):
                scored = [(max(s, 1), max(h, 1), d) for s, h, d in scored]
            else:
                scored = [(s + 1, h + 1, d) for s, h, d in scored]
        level_map = {}
        for s, _, d in scored:
            if d and s not in level_map:
                level_map[s] = d
        top = max(h for _, h, _ in scored)
        stated = _int(item.get("max_score"), 0)
        max_score = max(MIN_LEVELS, top, stated if top <= stated <= MAX_POINTS else 0)
        try:
            weight = float(item.get("weight") or 1)
        except (TypeError, ValueError):
            weight = 1.0
        key = re.sub(r"[^a-z0-9]+", "_", name.lower()).strip("_")
        if key in seen:
            warnings.append(f"“{name}” appeared twice in the file; the second one was kept as a separate row.")
        seen.add(key)
        criteria.append({"label": name, "max_score": max_score,
                         "weight": round(max(weight, 0), 2), "levels": level_map})
    # Weights stated as points or percentages (e.g. 25, 15): scale those so
    # their average is 1. Criteria without stated points keep the default 1.
    stated = [c["weight"] for c in criteria if c["weight"] > 5]
    if stated:
        mean = sum(stated) / len(stated)
        for c in criteria:
            if c["weight"] > 5:
                c["weight"] = round(c["weight"] / mean, 2)
    return criteria, warnings


def _words(label):
    return set(re.findall(r"[a-z0-9]+", label.lower())) - {"and", "the", "of", "a", "to", "in"}


def _add_missed(ai_criteria, table_criteria):
    """
    Criteria the table reader found that the AI left out (AIs sometimes stop
    early or skip rows). A table criterion counts as present when an AI
    criterion shares most of its words, e.g. "Thesis" vs "Thesis and Focus".
    """
    present = [_words(c["label"]) for c in ai_criteria]
    missed = []
    for c in table_criteria:
        words = _words(c["label"])
        if words and not any(len(words & p) >= max(1, min(len(words), len(p)) // 2 + 1) or words <= p or p <= words
                             for p in present if p):
            missed.append(c)
    return missed


def import_rubric(data, filename):
    suffix = os.path.splitext(filename or "")[1].lower()
    if suffix not in (".pdf", ".xlsx", ".csv"):
        raise ValueError("Upload a PDF, an Excel file (.xlsx) or a CSV.")

    tables = pdf_tables(data) if suffix == ".pdf" else sheet_tables(data, suffix)
    text = tables_as_text(tables)
    if suffix == ".pdf":
        text = (text + "\n\n" + pdf_text(data)).strip()

    attempts = []
    title = os.path.splitext(os.path.basename(filename))[0]
    readers = []
    if suffix == ".pdf":
        readers.append(("gemini", lambda: _gemini([
            {"inline_data": {"mime_type": "application/pdf", "data": base64.b64encode(data).decode()}},
            {"text": INSTRUCTIONS}])))
    else:
        readers.append(("gemini", lambda: _gemini([{"text": INSTRUCTIONS + "\n\nRubric (cells separated by |):\n\n" + text[:20000]}])))
    readers.append(("deepseek", lambda: _deepseek(text)))

    # Criteria from every table: some rubrics put one criterion on each page.
    found = [c for table in tables for c in table_rubric(table)]
    table_criteria, table_warnings = _normalize(found) if len(found) >= 2 else ([], [])

    for source, read in readers:
        try:
            result = read()
            criteria, warnings = _normalize(result.get("criteria") or [])
            if criteria:
                added = _add_missed(criteria, table_criteria)
                warnings += [f"“{c['label']}” was missed by {source.title()} and added from the file’s table; "
                             "please check it." for c in added]
                return {"title": _clean(result.get("title")) or title, "criteria": criteria + added,
                        "source": source, "warnings": warnings}
            attempts.append(f"{source}: no criteria found")
        except Exception as e:  # unavailable, quota, bad reply: try the next reader
            attempts.append(f"{source}: {str(e)[:120]}")
    print(f"[rubric import] AI readers unavailable or empty: {attempts}")

    criteria, warnings = table_criteria, table_warnings
    if criteria:
        return {"title": title, "criteria": criteria, "source": "table",
                "warnings": warnings + ["Read with the built-in table reader (no AI was available); "
                                        "please check the criteria."]}
    raise ValueError("No rubric criteria could be found in this file. "
                     "Check that it contains a grading grid, or add the criteria by hand.")
